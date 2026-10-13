"""Test suite for real tables: Markdown pipe -> Word Table / Telegram aligned <pre>."""
from docx import Document

from agent_bot.telegram.formatting import markdown_to_telegram_html
from agent_bot.tools.file_generator import FileGenerator
from agent_bot.tools.tables import (
    MdTable,
    align_pre,
    parse_markdown_tables,
    split_blocks,
)

MD = (
    "Kirish matni\n"
    "| Model | Narx |\n"
    "|---|---|\n"
    "| A | 10 |\n"
    "| B | 20 |\n"
    "Xulosa matni"
)


def test_parse_markdown_tables():
    tables = parse_markdown_tables(MD)
    assert len(tables) == 1
    assert tables[0].headers == ["Model", "Narx"]
    assert tables[0].rows == [["A", "10"], ["B", "20"]]


def test_parse_ignores_non_tables():
    assert parse_markdown_tables("a | b bitta qator\noddiy matn") == []
    assert parse_markdown_tables("") == []


def test_split_blocks():
    blocks = split_blocks(MD)
    kinds = [k for k, _ in blocks]
    assert kinds[0] == "text" and "table" in kinds and kinds[-1] == "text"
    table = next(p for k, p in blocks if k == "table")
    assert isinstance(table, MdTable) and table.headers == ["Model", "Narx"]


def test_align_pre_columns():
    out = align_pre(MdTable(headers=["Model", "Narx"], rows=[["A", "10"], ["B", "20"]]))
    lines = out.split("\n")
    assert len(lines) == 4  # header + separator + 2 rows
    assert all(len(ln) == len(lines[0]) for ln in lines)  # ustunlar tekis
    assert lines[1].startswith("|-")


def test_telegram_table_aligned():
    html = markdown_to_telegram_html(MD)
    assert "<pre>" in html and "Model" in html
    assert "|---|---|" not in html  # asl markdown separator tushirilgan
    assert "Kirish matni" in html and "Xulosa matni" in html


def test_docx_real_table(tmp_path):
    fg = FileGenerator(tmp_path)
    path = fg.generate_docx("Jadval testi", [{"title": "Narxlar", "content": MD}])
    doc = Document(str(path))
    assert len(doc.tables) == 1
    grid = doc.tables[0]
    assert grid.cell(0, 0).text == "Model" and grid.cell(1, 1).text == "10"
    assert len(grid.rows) == 3 and len(grid.columns) == 2
