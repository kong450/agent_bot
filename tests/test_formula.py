"""Test suite for DOCX formula rendering (Unicode math + super/subscript)."""
from docx import Document

from agent_bot.tools.file_generator import FileGenerator
from agent_bot.tools.formula import add_rich_paragraph, to_unicode_math


def test_unicode_math():
    assert "√" in to_unicode_math("sqrt(16) = 4")
    assert "π" in to_unicode_math("pi ~ 3.14")
    assert "→" in to_unicode_math("a -> b")
    assert "≤" in to_unicode_math("x <= y")
    # so'z ichi buzilmasligi kerak: 'spiral' ichidagi 'pi' tegilmaydi
    assert to_unicode_math("spiral") == "spiral"
    # sonlar orasidagi * ko'paytiruvga aylanadi
    assert "×" in to_unicode_math("3*4")


def test_superscript_subscript_runs():
    doc = Document()
    add_rich_paragraph(doc.add_paragraph(), "E = m*c^2, x_1 + a_{max}")
    runs = [r for p in doc.paragraphs for r in p.runs]
    assert any(r.text == "2" and r.font.superscript for r in runs)
    assert any(r.text == "1" and r.font.subscript for r in runs)
    assert any(r.text == "max" and r.font.subscript for r in runs)


def test_underscore_words_untouched():
    doc = Document()
    add_rich_paragraph(doc.add_paragraph(), "test_case saqlansin")
    runs = [r for p in doc.paragraphs for r in p.runs]
    assert not any(r.font.subscript for r in runs)
    assert "test_case" in "".join(r.text for r in runs)


def test_code_spans_untouched():
    doc = Document()
    add_rich_paragraph(doc.add_paragraph(), "kodda `a*b_c^d` o'zgarmaydi")
    runs = [r for p in doc.paragraphs for r in p.runs]
    mono = next(r for r in runs if r.font.name == "Consolas")
    assert mono.text == "a*b_c^d"
    assert not any(r.font.superscript or r.font.subscript for r in runs)


def test_variable_italic_math_font():
    doc = Document()
    add_rich_paragraph(doc.add_paragraph(), "E = m")
    runs = [r for p in doc.paragraphs for r in p.runs]
    e = next(r for r in runs if r.text == "E")
    assert e.font.italic and e.font.name == "Cambria Math"


def test_docx_end_to_end_with_formulas(tmp_path):
    fg = FileGenerator(tmp_path)
    path = fg.generate_docx(
        "Formula testi",
        [{"title": "Hisob", "content": "Quvvat P = U*I, kuchlanish U_1 = 220 V, maydon S = pi*r^2"}],
    )
    assert path.exists()
    doc = Document(str(path))
    runs = [r for p in doc.paragraphs for r in p.runs]
    texts = "".join(r.text for r in runs)
    assert "π" in texts and "×" in texts
    assert any(r.text == "1" and r.font.subscript for r in runs)
    assert any(r.text == "2" and r.font.superscript for r in runs)


def test_xlsx_rich_formula_cells(tmp_path):
    import zipfile

    from openpyxl import load_workbook

    fg = FileGenerator(tmp_path)
    path = fg.generate_xlsx("Smeta", ["Ko'rsatkich", "Qiymat"],
                            [["Quvvat", "P = U*I"], ["Maydon", "S = pi*r^2"], ["Son", 42]])
    # openpyxl reader rich-textni str ga yig'adi, shuning uchun xom XML tekshiriladi
    # (Excel'da chiroyli ko'rinadi — bu testlangan fakt)
    zf = zipfile.ZipFile(str(path))
    xml = zf.read("xl/worksheets/sheet1.xml").decode("utf-8")
    try:
        shared = zf.read("xl/sharedStrings.xml").decode("utf-8")
    except KeyError:
        shared = ""
    blob = xml + shared
    assert 'vertAlign val="superscript"' in blob  # r^2 darajasi
    # openpyxl ASCII-bo'lmagan belgilarni XML entity qiladi: × -> &#215;, π -> &#960;
    assert ("&#215;" in blob or "×" in blob) and ("&#960;" in blob or "π" in blob)
    wb = load_workbook(str(path))
    assert wb.active["B4"].value == 42  # sonlar sonligicha qoladi


def test_pptx_rich_formula_runs(tmp_path):
    from pptx import Presentation

    fg = FileGenerator(tmp_path)
    path = fg.generate_pptx("Taqdimot", [{"title": "Hisob",
                                          "points": ["P = U*I", "S = pi*r^2"]}])
    prs = Presentation(str(path))
    texts = []
    sup_found = sub_found = False
    for slide in prs.slides:
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            for para in shape.text_frame.paragraphs:
                for run in para.runs:
                    texts.append(run.text)
                    base = run._r.get_or_add_rPr().get("baseline")
                    if base == "30000":
                        sup_found = True
                    if base == "-25000":
                        sub_found = True
    joined = "".join(texts)
    assert "π" in joined and "×" in joined
    assert sup_found
