"""Professional Document, Spreadsheet, and Presentation Generator."""
from pathlib import Path
from typing import List, Dict, Any, Optional
import datetime

class FileGenerator:
    """DOCX, XLSX va PPTX fayllarini professional ko'rinishda yaratuvchi asbob."""

    def __init__(self, output_dir: Path):
        self.output_dir = output_dir
        self.output_dir.mkdir(parents=True, exist_ok=True)

    def generate_docx(
        self,
        title: str,
        sections: List[Dict[str, str]],
        filename: Optional[str] = None
    ) -> Path:
        """Professional Word (DOCX) hisobotini generatsiya qiladi."""
        from docx import Document
        from docx.shared import Pt, Inches, RGBColor
        from docx.enum.text import WD_ALIGN_PARAGRAPH

        doc = Document()

        # Sarlavha
        title_p = doc.add_paragraph()
        title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        title_run = title_p.add_run(f"⚡ {title}\n")
        title_run.font.size = Pt(20)
        title_run.font.bold = True
        title_run.font.color.rgb = RGBColor(16, 44, 87)

        # Sana va Agent Bot belgisi
        meta_p = doc.add_paragraph()
        meta_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        now_str = datetime.datetime.now().strftime("%Y-%m-%d %H:%M")
        meta_run = meta_p.add_run(f"Agent Bot AI Platform | Sana: {now_str}\n")
        meta_run.font.size = Pt(10)
        meta_run.font.italic = True

        doc.add_paragraph("―" * 45)

        # Bo'limlar
        for sec in sections:
            sec_title = sec.get("title", "")
            sec_content = sec.get("content", "")

            if sec_title:
                h = doc.add_heading(level=1)
                h_run = h.add_run(sec_title)
                h_run.font.size = Pt(14)
                h_run.font.bold = True

            if sec_content:
                from agent_bot.tools.formula import add_rich_paragraph
                from agent_bot.tools.tables import add_word_table, split_blocks

                for kind, payload in split_blocks(sec_content):
                    if kind == "table":
                        add_word_table(doc, payload)
                        continue
                    for block in str(payload).split("\n"):
                        if not block.strip():
                            continue
                        p = doc.add_paragraph()
                        p.paragraph_format.line_spacing = 1.15
                        p.paragraph_format.space_after = Pt(4)
                        add_rich_paragraph(p, block)

        if not filename:
            safe_title = "".join(c for c in title if c.isalnum() or c in (' ', '_', '-')).strip()
            filename = f"{safe_title.replace(' ', '_')}_{int(datetime.datetime.now().timestamp())}.docx"

        file_path = self.output_dir / filename
        doc.save(str(file_path))
        return file_path

    def generate_xlsx(
        self,
        sheet_name: str,
        headers: List[str],
        rows: List[List[Any]],
        filename: Optional[str] = None
    ) -> Path:
        """Excel (XLSX) hisob-kitob smetasi yoki jadvalini generatsiya qiladi."""
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

        wb = Workbook()
        ws = wb.active
        ws.title = sheet_name

        # Header dizayni (To'q ko'k rang va oq yozuv)
        header_fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        center_align = Alignment(horizontal="center", vertical="center")
        thin_border = Border(
            left=Side(style='thin', color='D9D9D9'),
            right=Side(style='thin', color='D9D9D9'),
            top=Side(style='thin', color='D9D9D9'),
            bottom=Side(style='thin', color='D9D9D9')
        )

        ws.append(headers)
        for col_num in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col_num)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = center_align

        # Qatorlar (matnli kataklar formula-formatda: x^2, x_1, sqrt, pi)
        from agent_bot.tools.formula import to_cell_rich_text

        for row in rows:
            ws.append([
                to_cell_rich_text(str(v)) if isinstance(v, str) else v
                for v in row
            ])

        # Kataklar chegarasi va kengligini sozlash
        for row in ws.iter_rows(min_row=1, max_row=len(rows) + 1, min_col=1, max_col=len(headers)):
            for cell in row:
                cell.border = thin_border

        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = col[0].column_letter
            ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

        if not filename:
            filename = f"Smeta_{int(datetime.datetime.now().timestamp())}.xlsx"

        file_path = self.output_dir / filename
        wb.save(str(file_path))
        return file_path

    def generate_pptx(
        self,
        title: str,
        slides_data: List[Dict[str, Any]],
        filename: Optional[str] = None
    ) -> Path:
        """PowerPoint (PPTX) taqdimot slaydlarini yaratadi."""
        from pptx import Presentation
        from pptx.util import Inches, Pt

        prs = Presentation()

        # 1-Slayd: Muqova
        title_slide_layout = prs.slide_layouts[0]
        slide = prs.slides.add_slide(title_slide_layout)
        title_box = slide.shapes.title
        subtitle_box = slide.placeholders[1]

        title_box.text = title
        subtitle_box.text = f"Agent Bot AI | {datetime.datetime.now().strftime('%Y-%m-%d')}"

        # Keyingi slaydlar: Kontent (formulalar baseline formatda)
        from agent_bot.tools.formula import add_pptx_runs

        bullet_layout = prs.slide_layouts[1]
        for s_data in slides_data:
            slide = prs.slides.add_slide(bullet_layout)
            slide.shapes.title.text = s_data.get("title", "")
            tf = slide.placeholders[1].text_frame
            tf.clear()

            points = s_data.get("points", [])
            for idx, pt in enumerate(points):
                p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
                p.level = 0
                add_pptx_runs(p, str(pt))

        if not filename:
            safe_title = "".join(c for c in title if c.isalnum() or c in (' ', '_', '-')).strip()
            filename = f"{safe_title.replace(' ', '_')}_{int(datetime.datetime.now().timestamp())}.pptx"

        file_path = self.output_dir / filename
        prs.save(str(file_path))
        return file_path
