"""Universal document ingestion and text extraction utility (PDF, DOCX, TXT, MD, CSV)."""
from pathlib import Path
from typing import Dict, Any, Optional

class FileReader:
    """Foydalanuvchi yuborgan hujjatlardan matn va ma'lumotlarni xavfsiz o'qib beruvchi modul."""

    @classmethod
    def read_file(cls, file_path: Path) -> Dict[str, Any]:
        """Fayl turiga qarab matnni ajratib oladi."""
        path = Path(file_path)
        if not path.exists():
            return {"success": False, "error": "Fayl topilmadi.", "text": ""}

        suffix = path.suffix.lower()

        try:
            if suffix in [".txt", ".md", ".csv", ".json", ".py"]:
                text = path.read_text(encoding="utf-8", errors="replace")
                return {"success": True, "text": text, "format": suffix}

            elif suffix == ".docx":
                return cls._read_docx(path)

            elif suffix == ".pdf":
                return cls._read_pdf(path)

            else:
                return {
                    "success": False,
                    "error": f"Ushbu fayl formati ({suffix}) hali qo'llab-quvvatlanmaydi.",
                    "text": ""
                }
        except Exception as e:
            return {"success": False, "error": f"Faylni o'qishda xatolik: {str(e)}", "text": ""}

    @classmethod
    def _read_docx(cls, path: Path) -> Dict[str, Any]:
        """Word (DOCX) hujjatidan matnni ajratadi."""
        from docx import Document
        doc = Document(str(path))
        full_text = []
        for para in doc.paragraphs:
            if para.text.strip():
                full_text.append(para.text)

        # Jadvallardagi matnlarni ham o'qish
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                if row_text:
                    full_text.append(row_text)

        return {
            "success": True,
            "text": "\n".join(full_text),
            "format": ".docx",
            "paragraphs_count": len(doc.paragraphs)
        }

    @classmethod
    def _read_pdf(cls, path: Path) -> Dict[str, Any]:
        """PDF hujjatidan matnni ajratadi (agar pypdf mavjud bo'lsa)."""
        try:
            import pypdf
            reader = pypdf.PdfReader(str(path))
            pages_text = []
            for idx, page in enumerate(reader.pages):
                extracted = page.extract_text()
                if extracted:
                    pages_text.append(f"--- [Sahifa {idx + 1}] ---\n{extracted.strip()}")

            return {
                "success": True,
                "text": "\n\n".join(pages_text),
                "format": ".pdf",
                "pages_count": len(reader.pages)
            }
        except ImportError:
            # Fallback agar pypdf o'rnatilmagan bo'lsa
            return {
                "success": False,
                "error": "PDF o'qish uchun 'pypdf' kutubxonasi o'rnatilmagan. Matnli fayllardan foydalaning.",
                "text": ""
            }
